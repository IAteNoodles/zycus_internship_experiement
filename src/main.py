from __future__ import annotations
import argparse
import os
from datetime import date
from typing import Optional

from dotenv import load_dotenv

from database import parse_date
from data import sample_projects
from models.projects import Project, SentimentScore
from models.rag import compute_rag, RagResult, RAG
from models.sentiment import analyze_sentiment
from reports import weekly_narrative, monthly_content

load_dotenv()


def _enrich_sentiment(project: Project) -> None:
    snap = project.latest
    if not snap:
        return
    for entry in snap.stakeholder_sentiment:
        entry.score = analyze_sentiment(entry.comment)


ANALYSIS_HEADER = "=" * 60


def run_weekly(projects: list[Project], as_of: date) -> list[tuple[Project, RagResult]]:
    results = []
    for proj in projects:
        if not proj.latest:
            print(f"[SKIP] {proj.name}: no snapshot.")
            continue
        _enrich_sentiment(proj)
        result = compute_rag(proj, proj.latest, as_of)
        results.append((proj, result))
        print(ANALYSIS_HEADER)
        print(weekly_narrative(proj, result))
        print()
    return results


def synthesize_monthly(results: list[tuple[Project, RagResult]], as_of: date, path: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    content = monthly_content(results, as_of)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W = prs.slide_width
    H = prs.slide_height

    MARGIN_L = Inches(0.7)
    MARGIN_T = Inches(0.5)
    CONTENT_W = W - Inches(1.4)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK = RGBColor(0x1E, 0x29, 0x3B)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)
    BORDER_LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
    RAG_COLORS = {"Green": RGBColor(0x22, 0xC5, 0x5E), "Amber": RGBColor(0xF5, 0x9E, 0x0B), "Red": RGBColor(0xEF, 0x44, 0x44)}
    SIGNAL_COLORS = {"schedule": RGBColor(0x3B, 0x82, 0xF6), "budget": RGBColor(0x8B, 0x5C, 0xF6), "blockers": RGBColor(0xF5, 0x9E, 0x0B), "sentiment": RGBColor(0x06, 0xB6, 0xD4)}

    def _tb(sl, left, top, width, height):
        tf = sl.shapes.add_textbox(left, top, width, height).text_frame
        tf.word_wrap = True
        return tf

    def _p(tf, text, size=Pt(16), color=DARK, bold=False, align=PP_ALIGN.LEFT, space_after=Pt(6)):
        if not tf.paragraphs[0].text:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = size
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = align
        p.space_after = space_after
        return p

    def _rect(sl, left, top, w, h, fill, line=None):
        shape = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        if line:
            shape.line.color.rgb = line
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def _rrect(sl, left, top, w, h, fill, line=None, text="", text_color=WHITE, font_size=Pt(14), bold=True, align=PP_ALIGN.CENTER):
        shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        if line:
            shape.line.color.rgb = line
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        if text:
            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = align
            p.text = text
            p.font.size = font_size
            p.font.bold = bold
            p.font.color.rgb = text_color
        return shape

    def _rag_badge(sl, status, left, top, size=Inches(1.0)):
        return _rrect(sl, left, top, size, Inches(0.5), RAG_COLORS.get(status, MUTED), text=status.upper(), text_color=WHITE, font_size=Pt(16), bold=True)

    def _score_bar(sl, score, left, top, w, h, color):
        bg = _rrect(sl, left, top, w, h, BORDER_LIGHT)
        if score is not None and score > 0:
            fill_w = int(w * (score / 2.0))
            _rrect(sl, left, top, fill_w, h, color)
        label = f"{score}/2" if score is not None else "N/A"
        tf = _tb(sl, left + Inches(0.15), top, w - Inches(0.3), h)
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE if (score is not None and score > 0) else DARK

    # ── SLIDE 1: Title ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = DARK
    _rect(sl, Inches(0), Inches(3.0), W, Inches(0.06), ACCENT)
    tf = _tb(sl, Inches(0), Inches(1.5), W, Inches(1.5))
    _p(tf, "Monthly Project Health Report", Pt(44), WHITE, True, PP_ALIGN.CENTER, Pt(8))
    tf2 = _tb(sl, Inches(0), Inches(3.3), W, Inches(0.8))
    _p(tf2, as_of.strftime("%B %d, %Y"), Pt(22), RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.CENTER)
    tf3 = _tb(sl, Inches(0), Inches(4.5), W, Inches(0.6))
    _p(tf3, f"{len(results)} projects evaluated", Pt(16), MUTED, align=PP_ALIGN.CENTER)

    # ── SLIDE 2: Executive Summary + RAG Donut ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT_BG
    _rrect(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6), DARK, text="Executive Summary", text_color=WHITE, font_size=Pt(22), bold=True)

    _p(_tb(sl, MARGIN_L, Inches(1.5), Inches(5.5), Inches(3.0)),
       content.get("executive_summary", "No executive summary available."), Pt(15), DARK)

    green = sum(1 for _, r in results if r.status == RAG.GREEN)
    amber = sum(1 for _, r in results if r.status == RAG.AMBER)
    red = sum(1 for _, r in results if r.status == RAG.RED)
    total = len(results)

    cd = CategoryChartData()
    cd.categories = ["Green", "Amber", "Red"]
    cd.add_series("Projects", (green, amber, red))
    cf = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(7.0), Inches(1.3), Inches(5.0), Inches(4.0), cd)
    chart = cf.chart
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    plot = chart.plots[0]
    for i, c in enumerate([RAG_COLORS["Green"], RAG_COLORS["Amber"], RAG_COLORS["Red"]]):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    chart.has_title = False

    y = Inches(5.5)
    cw = Inches(3.8)
    gap = Inches(0.3)
    sx = MARGIN_L
    for label, count, color in [("Total Projects", total, ACCENT), ("On Track", green, RAG_COLORS["Green"]), ("At Risk", amber + red, RAG_COLORS["Red"])]:
        card = _rrect(sl, sx, y, cw, Inches(1.2), CARD_BG, BORDER_LIGHT)
        _rect(sl, sx, y, cw, Inches(0.06), color)
        _p(_tb(sl, sx, y + Inches(0.15), cw, Inches(0.7)), str(count), Pt(36), color, True, PP_ALIGN.CENTER, Pt(0))
        _p(_tb(sl, sx, y + Inches(0.7), cw, Inches(0.4)), label, Pt(13), MUTED, align=PP_ALIGN.CENTER)
        sx += cw + gap

    # ── SLIDES 3+: Project Detail with Score Bars ──
    for proj, result in results:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT_BG

        _rrect(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6), DARK, text=proj.name, text_color=WHITE, font_size=Pt(22), bold=True)
        _rag_badge(sl, result.status, MARGIN_L + CONTENT_W - Inches(1.0), MARGIN_T)

        meta = f"Budget: ${proj.budget:,.0f}  |  {proj.start_date} \u2192 {proj.end_date}"
        if result.weighted_score is not None:
            meta += f"  |  Score: {result.weighted_score:.2f}"
        tfm = _tb(sl, MARGIN_L, Inches(1.3), CONTENT_W, Inches(0.3))
        _p(tfm, meta, Pt(11), MUTED, space_after=Pt(0))

        gl = MARGIN_L
        gt = Inches(1.8)
        bw = Inches(4.6)
        bh = Inches(0.38)
        lw = Inches(1.1)
        gx = Inches(0.5)
        gy = Inches(0.7)

        for i, sig in enumerate(result.signals):
            col = i % 2
            row = i // 2
            x = gl + col * (lw + bw + gx)
            y = gt + row * (bh + gy)

            lt = _tb(sl, x, y, lw, bh)
            _p(lt, sig.name.capitalize(), Pt(13), SIGNAL_COLORS.get(sig.name, DARK), True, space_after=Pt(0))
            lt.margin_left = Inches(0)
            lt.margin_top = Inches(0.06)

            _score_bar(sl, sig.score, x + lw, y, bw, bh, SIGNAL_COLORS.get(sig.name, ACCENT))

            rt = _tb(sl, x + lw, y + bh + Inches(0.02), bw, Inches(0.3))
            _p(rt, sig.rationale, Pt(8), MUTED, space_after=Pt(0))

        yo = gt + 2 * (bh + gy) + Inches(0.2)
        if result.overrides_applied:
            _p(_tb(sl, MARGIN_L, yo, CONTENT_W, Inches(0.4)),
               "\u26a0 " + "; ".join(result.overrides_applied), Pt(11), RGBColor(0xD9, 0x77, 0x06))
            yo += Inches(0.4)

        narrative = weekly_narrative(proj, result)
        ny = max(yo + Inches(0.1), Inches(4.8))
        remaining = H - ny - Inches(0.5)
        if remaining > Inches(0.5):
            ch = min(Inches(2.0), remaining)
            _rrect(sl, MARGIN_L, ny, CONTENT_W, ch, CARD_BG, BORDER_LIGHT)
            _p(_tb(sl, MARGIN_L + Inches(0.25), ny + Inches(0.15), CONTENT_W - Inches(0.5), ch - Inches(0.3)),
               narrative, Pt(12), DARK)

    # ── TRENDS & RISKS ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT_BG
    _rrect(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6), DARK, text="Trends & Emerging Risks", text_color=WHITE, font_size=Pt(22), bold=True)

    trends = content.get("trends", [])
    risks = content.get("risks", [])
    cw2 = Inches(5.8)
    sep_x = MARGIN_L + cw2 + Inches(0.25)
    rt = _tb(sl, MARGIN_L, Inches(1.5), cw2, Inches(4.5))
    _p(rt, "Trends", Pt(20), SIGNAL_COLORS["schedule"], True, space_after=Pt(12))
    if trends:
        for t in trends:
            _p(rt, f"  \u25b8 {t}", Pt(14), DARK, space_after=Pt(8))
    else:
        _p(rt, "No significant trends detected.", Pt(14), MUTED)
    rt2 = _tb(sl, sep_x + Inches(0.35), Inches(1.5), cw2, Inches(4.5))
    _p(rt2, "Risks", Pt(20), RAG_COLORS["Red"], True, space_after=Pt(12))
    if risks:
        for r in risks:
            _p(rt2, f"  \u25b8 {r}", Pt(14), DARK, space_after=Pt(8))
    else:
        _p(rt2, "No significant risks identified.", Pt(14), MUTED)
    _rect(sl, sep_x, Inches(1.5), Inches(0.02), Inches(4.5), BORDER_LIGHT)

    # ── RECOMMENDATIONS ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.background; bg.fill.solid(); bg.fill.fore_color.rgb = LIGHT_BG
    _rrect(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6), DARK, text="Recommendations", text_color=WHITE, font_size=Pt(22), bold=True)

    recs = content.get("recommendations", [])
    if recs:
        for i, rec in enumerate(recs, 1):
            cy = Inches(1.5) + (i - 1) * Inches(0.9)
            _rrect(sl, MARGIN_L, cy, CONTENT_W, Inches(0.75), CARD_BG, BORDER_LIGHT)
            circle = sl.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN_L + Inches(0.2), cy + Inches(0.15), Inches(0.45), Inches(0.45))
            circle.fill.solid()
            circle.fill.fore_color.rgb = ACCENT
            circle.line.fill.background()
            cp = circle.text_frame.paragraphs[0]
            cp.text = str(i)
            cp.font.size = Pt(16)
            cp.font.bold = True
            cp.font.color.rgb = WHITE
            cp.alignment = PP_ALIGN.CENTER
            _p(_tb(sl, MARGIN_L + Inches(0.85), cy + Inches(0.12), CONTENT_W - Inches(1.2), Inches(0.55)),
               rec, Pt(14), DARK, space_after=Pt(0))
    else:
        _p(_tb(sl, MARGIN_L, Inches(1.5), CONTENT_W, Inches(1.0)),
           "No recommendations at this time.", Pt(16), MUTED)

    prs.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description="Project Health Reporting Agent")
    parser.add_argument("--output-dir", default="output")
    parser.add_argument("--date", default=None, help="YYYY-MM-DD")
    parser.add_argument("--synthesize", action="store_true", help="Generate monthly PPTX")
    args = parser.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)
    as_of = parse_date(args.date) if args.date else date.today()

    projects = sample_projects()
    results = run_weekly(projects, as_of)

    weekly_path = os.path.join(args.output_dir, f"weekly_{as_of.isoformat()}.txt")
    with open(weekly_path, "w") as f:
        for proj, result in results:
            f.write(weekly_narrative(proj, result) + "\n\n")
    print(f"Weekly: {weekly_path}")

    if args.synthesize:
        pptx = synthesize_monthly(results, as_of, os.path.join(args.output_dir, f"monthly_{as_of.isoformat()}.pptx"))
        print(f"Monthly: {pptx}")


if __name__ == "__main__":
    main()
