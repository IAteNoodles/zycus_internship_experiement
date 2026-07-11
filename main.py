from __future__ import annotations
import argparse
import os
from datetime import date
from typing import Optional

from dotenv import load_dotenv

from data import sample_projects
from projects import Project, SentimentScore
from rag import compute_rag, RagResult, RAG
from sentiment import analyze_sentiment
from reports import weekly_narrative, monthly_content

load_dotenv()


def _enrich_sentiment(project: Project) -> None:
    snap = project.latest
    if not snap:
        return
    for entry in snap.stakeholder_sentiment:
        entry.score = analyze_sentiment(entry.comment)


ANALYSIS_HEADER = "=" * 60


def run_weekly(projects: list[Project], as_of: date) -> list[tuple[Project, RagResult]]:
    results = []
    for proj in projects:
        if not proj.latest:
            print(f"[SKIP] {proj.name}: no snapshot.")
            continue
        _enrich_sentiment(proj)
        result = compute_rag(proj, proj.latest, as_of)
        results.append((proj, result))
        print(ANALYSIS_HEADER)
        print(weekly_narrative(proj, result))
        print()
    return results


def synthesize_monthly(results: list[tuple[Project, RagResult]], as_of: date, path: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    content = monthly_content(results, as_of)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W = prs.slide_width

    MARGIN_L = Inches(0.7)
    MARGIN_T = Inches(0.5)
    CONTENT_W = W - Inches(1.4)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK = RGBColor(0x1E, 0x29, 0x3B)
    MUTED = RGBColor(0x64, 0x74, 0x8B)
    RAG_COLORS = {"Green": RGBColor(0x22, 0xC5, 0x5E), "Amber": RGBColor(0xF5, 0x9E, 0x0B), "Red": RGBColor(0xEF, 0x44, 0x44)}

    def _add_bg(sl, color=DARK):
        bg = sl.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def _tb(sl, left, top, width, height):
        return sl.shapes.add_textbox(left, top, width, height).text_frame

    def _add_para(tf, text, size=Pt(16), color=WHITE, bold=False, align=PP_ALIGN.LEFT, space_after=Pt(6)):
        p = tf.add_paragraph()
        p.text = text
        p.font.size = size
        p.font.color.rgb = color
        p.font.bold = bold
        p.alignment = align
        p.space_after = space_after
        return p

    def _rag_badge(sl, status, left, top, size=Inches(0.9)):
        shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, size, Inches(0.45))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RAG_COLORS.get(status, MUTED)
        shape.line.fill.background()
        tf = shape.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].text = status.upper()
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        return shape

    def _signal_table(sl, signals, top, left=MARGIN_L, width=CONTENT_W):
        rows, cols = len(signals) + 1, 3
        table = sl.shapes.add_table(rows, cols, left, top, width, Inches(0.35 * rows)).table
        for i, h in enumerate(["Signal", "Score", "Rationale"]):
            c = table.cell(0, i)
            c.text = h
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.bold = True
                p.font.color.rgb = DARK
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        for ri, s in enumerate(signals, 1):
            for ci, val in enumerate([s.name, f"{s.score}/2" if s.score is not None else "N/A", s.rationale]):
                c = table.cell(ri, ci)
                c.text = val
                for p in c.text_frame.paragraphs:
                    p.font.size = Pt(11)
                    p.font.color.rgb = DARK if ci == 2 else RAG_COLORS.get({"schedule": "Green", "budget": "Green", "blockers": "Green", "sentiment": "Green"}.get(s.name, "Amber"), DARK)
                    if ci == 2:
                        p.font.color.rgb = MUTED
                c.fill.solid()
                c.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF8, 0xFA, 0xFC)
        return table

    # ── Title Slide ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(sl, DARK)
    tf = _tb(sl, Inches(0), Inches(2.0), W, Inches(1.5))
    _add_para(tf, "Monthly Project Health Report", Pt(40), WHITE, True, PP_ALIGN.CENTER, Pt(0))
    tf2 = _tb(sl, Inches(0), Inches(3.5), W, Inches(1))
    _add_para(tf2, as_of.strftime(f"%B {as_of.day}, %Y"), Pt(20), RGBColor(0x94, 0xA3, 0xB8), False, PP_ALIGN.CENTER)

    # ── Executive Summary ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(sl)
    tf = _tb(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6))
    _add_para(tf, "Executive Summary", Pt(28), DARK, True)
    tf2 = _tb(sl, MARGIN_L, Inches(1.3), CONTENT_W, Inches(2.5))
    _add_para(tf2, content.get("executive_summary", "No executive summary available."), Pt(16), MUTED)

    green = sum(1 for _, r in results if r.status == RAG.GREEN)
    amber = sum(1 for _, r in results if r.status == RAG.AMBER)
    red = sum(1 for _, r in results if r.status == RAG.RED)
    y = Inches(4.0)
    for status, count, color in [("Green", green, RAG_COLORS["Green"]), ("Amber", amber, RAG_COLORS["Amber"]), ("Red", red, RAG_COLORS["Red"])]:
        shape = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN_L + Inches(2.5 * ["Green", "Amber", "Red"].index(status)), y, Inches(2.0), Inches(0.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        tf = shape.text_frame
        tf.paragraphs[0].text = f"{status}: {count}"
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # ── Per-Project Slides ──
    for proj, result in results:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        _add_bg(sl)
        _rag_badge(sl, result.status, MARGIN_L, MARGIN_T)
        tf = _tb(sl, MARGIN_L + Inches(1.2), MARGIN_T, CONTENT_W - Inches(1.2), Inches(0.6))
        _add_para(tf, proj.name, Pt(26), DARK, True)
        _add_para(tf, f"Budget: ${proj.budget:,.0f}  |  {proj.start_date} → {proj.end_date}", Pt(12), MUTED)

        _signal_table(sl, result.signals, Inches(1.6))

        y = Inches(1.6) + Inches(0.35 * (len(result.signals) + 1)) + Inches(0.15)
        if result.overrides_applied:
            tf = _tb(sl, MARGIN_L, y, CONTENT_W, Inches(0.4))
            _add_para(tf, "Overrides: " + "; ".join(result.overrides_applied), Pt(11), RGBColor(0xD9, 0x77, 0x06))
            y += Inches(0.4)

        narrative = weekly_narrative(proj, result)
        remaining = prs.slide_height - y - Inches(0.5)
        if remaining > Inches(0.5):
            tf = _tb(sl, MARGIN_L, y, CONTENT_W, max(Inches(1.0), remaining))
            _add_para(tf, narrative, Pt(13), MUTED)

    # ── Trends ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(sl)
    tf = _tb(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6))
    _add_para(tf, "Trends & Emerging Risks", Pt(28), DARK, True)
    tf2 = _tb(sl, MARGIN_L, Inches(1.3), CONTENT_W, Inches(4.0))
    trends = content.get("trends", [])
    risks = content.get("risks", [])
    if trends:
        _add_para(tf2, "Trends", Pt(18), DARK, True)
        for t in trends:
            _add_para(tf2, f"  • {t}", Pt(14), MUTED)
    if risks:
        _add_para(tf2, "Risks", Pt(18), DARK, True, space_after=Pt(4))
        for r in risks:
            _add_para(tf2, f"  • {r}", Pt(14), MUTED)
    if not trends and not risks:
        _add_para(tf2, "No significant trends detected.", Pt(14), MUTED)

    # ── Recommendations ──
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(sl)
    tf = _tb(sl, MARGIN_L, MARGIN_T, CONTENT_W, Inches(0.6))
    _add_para(tf, "Recommendations", Pt(28), DARK, True)
    tf2 = _tb(sl, MARGIN_L, Inches(1.3), CONTENT_W, Inches(4.0))
    recs = content.get("recommendations", [])
    if recs:
        for rec in recs:
            _add_para(tf2, f"  ➤ {rec}", Pt(14), DARK, space_after=Pt(10))
    else:
        _add_para(tf2, "No recommendations at this time.", Pt(14), MUTED)

    prs.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description="Project Health Reporting Agent")
    parser.add_argument("--output-dir", default="output")
    parser.add_argument("--date", default=None, help="YYYY-MM-DD")
    parser.add_argument("--synthesize", action="store_true", help="Generate monthly PPTX")
    args = parser.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)
    as_of = date.fromisoformat(args.date) if args.date else date.today()

    projects = sample_projects()
    results = run_weekly(projects, as_of)

    weekly_path = os.path.join(args.output_dir, f"weekly_{as_of.isoformat()}.txt")
    with open(weekly_path, "w") as f:
        for proj, result in results:
            f.write(weekly_narrative(proj, result) + "\n\n")
    print(f"Weekly: {weekly_path}")

    if args.synthesize:
        pptx = synthesize_monthly(results, as_of, os.path.join(args.output_dir, f"monthly_{as_of.isoformat()}.pptx"))
        print(f"Monthly: {pptx}")


if __name__ == "__main__":
    main()
